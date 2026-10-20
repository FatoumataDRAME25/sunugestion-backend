from rest_framework import serializers
from django.contrib.auth import get_user_model
import re

from authentication.validators import valider_telephone

Utilisateur = get_user_model()

ROLES_AUTORISES = [
    'president',
    'tresorier',
    'secretaire',
    'membre'
]


def valider_role(role):
    """Vérifie que le rôle est autorisé."""
    role = str(role).strip().lower()

    if role not in ROLES_AUTORISES:
        return role, [
            f"Rôle invalide — valeurs autorisées : "
            f"{', '.join(ROLES_AUTORISES)}"
        ]

    return role, []


# Fonctions du bureau occupées par une seule personne à la fois dans un GIE
# (le président est géré à part, par la passation de présidence)
FONCTIONS_UNIQUES = {
    'tresorier': 'trésorier',
    'secretaire': 'secrétaire',
}


def erreur_role_nouveau_membre(gie, role, roles_deja_pris=()):
    """
    Rôle d'un membre ajouté ou importé :
    - président : jamais, la présidence se transmet par passation ;
    - trésorier, secrétaire : seulement si le poste est libre.
    `roles_deja_pris` : fonctions déjà attribuées plus haut dans le même fichier.
    Retourne le message d'erreur, ou None si le rôle est accepté.
    """
    if role == 'president':
        return (
            "Le rôle président ne peut pas être attribué à un nouveau membre. "
            "La présidence se transmet depuis la fiche d'un membre."
        )

    if role in FONCTIONS_UNIQUES:
        libelle = FONCTIONS_UNIQUES[role]
        if role in roles_deja_pris:
            return f"Le fichier contient plusieurs {libelle}s : un seul est possible."

        occupant = Utilisateur.objects.filter(gie=gie, role=role).first()
        if occupant:
            return (
                f"Le GIE a déjà un {libelle} ({occupant.prenom} {occupant.nom}). "
                f"Ajoutez ce membre comme membre simple."
            )

    return None


def valider_nom_prenom(valeur, champ):
    """Vérifie qu'un nom ou prénom contient uniquement des caractères valides."""
    valeur = str(valeur).strip()

    if not valeur:
        raise serializers.ValidationError(
            f"Le {champ} ne peut pas être vide."
        )

    if not re.match(r"^[A-Za-zÀ-ÖØ-öø-ÿ\s'-]+$", valeur):
        raise serializers.ValidationError(
            f"Le {champ} doit contenir uniquement des lettres."
        )

    return valeur


# ==========================================================
# AJOUT MANUEL
# ==========================================================

class AjoutMembreSerializer(serializers.Serializer):

    prenom = serializers.CharField(max_length=100)
    nom = serializers.CharField(max_length=100)
    telephone = serializers.CharField(max_length=20)
    role = serializers.ChoiceField(choices=ROLES_AUTORISES)

    def validate_telephone(self, valeur):
        telephone, erreurs = valider_telephone(valeur)
        if erreurs:
            raise serializers.ValidationError(erreurs)
        return telephone

    def validate_prenom(self, valeur):
        return valider_nom_prenom(valeur, "prénom")

    def validate_nom(self, valeur):
        return valider_nom_prenom(valeur, "nom")

    def validate_role(self, valeur):
        erreur = erreur_role_nouveau_membre(
            self.context['gie'], valeur, self.context.get('roles_deja_pris', ())
        )
        if erreur:
            raise serializers.ValidationError(erreur)
        return valeur

    def create(self, validated_data):

        gie = self.context['gie']

        membre = Utilisateur.objects.create_user(
            telephone=validated_data['telephone'],
            nom=validated_data['nom'],
            prenom=validated_data['prenom'],
            role=validated_data['role'],
            gie=gie,
            statut='en_attente'
        )

        membre.statut = 'en_attente'
        membre.save()

        # Génération du token d'invitation
        membre.generer_token_invitation()

        # Envoi du lien d'activation par SMS
        from membres.utils import envoyer_invitation
        envoyer_invitation(membre)

        return membre


# ==========================================================
# IMPORT EXCEL
# ==========================================================

class ImportExcelSerializer(serializers.Serializer):

    fichier = serializers.FileField()

    TAILLE_MAX_OCTETS = 5 * 1024 * 1024  # 5 Mo

    def validate_fichier(self, fichier):

        if not fichier.name.endswith('.xlsx'):
            raise serializers.ValidationError(
                "Le fichier doit être au format .xlsx"
            )

        if fichier.size > self.TAILLE_MAX_OCTETS:
            raise serializers.ValidationError(
                "Le fichier est trop volumineux (5 Mo maximum)."
            )

        return fichier

    def analyser(self):

        import openpyxl

        fichier = self.validated_data['fichier']

        # Un fichier renommé en .xlsx ou endommagé fait échouer openpyxl
        try:
            wb = openpyxl.load_workbook(fichier)
        except Exception:
            raise serializers.ValidationError(
                "Le fichier est illisible ou n'est pas un vrai fichier Excel (.xlsx)."
            )
        ws = wb.active

        colonnes_requises = ['prenom', 'nom', 'telephone', 'role']

        membres_valides = []
        membres_invalides = []
        gie = self.context.get('gie')
        roles_du_fichier = set()

        # Récupération des colonnes
        headers = [
            str(cell.value).strip().lower()
            for cell in ws[1]
        ]

        # Vérification des colonnes obligatoires
        for col in colonnes_requises:
            if col not in headers:
                raise serializers.ValidationError(
                    f"Colonne manquante dans le fichier : '{col}'"
                )

        # Lecture des lignes
        for numero_ligne, ligne in enumerate(
            ws.iter_rows(min_row=2, values_only=True),
            start=2
        ):
            # Ignorer les lignes complètement vides
            if not any(ligne):
                continue

            donnees = dict(zip(headers, ligne))
            erreurs_ligne = []

            prenom = str(donnees.get('prenom', '') or '').strip()
            nom = str(donnees.get('nom', '') or '').strip()
            telephone = str(donnees.get('telephone', '') or '').strip()
            role = str(donnees.get('role', '') or '').strip().lower()

            # Prénom
            if not prenom:
                erreurs_ligne.append("Prénom manquant")
            else:
                try:
                    prenom = valider_nom_prenom(prenom, "prénom")
                except serializers.ValidationError as erreur:
                    erreurs_ligne.extend(erreur.detail)

            # Nom
            if not nom:
                erreurs_ligne.append("Nom manquant")
            else:
                try:
                    nom = valider_nom_prenom(nom, "nom")
                except serializers.ValidationError as erreur:
                    erreurs_ligne.extend(erreur.detail)

            # Téléphone
            if not telephone:
                erreurs_ligne.append("Téléphone manquant")
            else:
                telephone, erreurs_tel = valider_telephone(telephone)
                erreurs_ligne.extend(erreurs_tel)

            # Rôle
            if not role:
                erreurs_ligne.append("Rôle manquant")
            else:
                role, erreurs_role = valider_role(role)
                erreurs_ligne.extend(erreurs_role)

                if not erreurs_role and gie is not None:
                    erreur = erreur_role_nouveau_membre(gie, role, roles_du_fichier)
                    if erreur:
                        erreurs_ligne.append(erreur)
                    elif role in FONCTIONS_UNIQUES:
                        roles_du_fichier.add(role)

            membre = {
                'ligne': numero_ligne,
                'prenom': prenom,
                'nom': nom,
                'telephone': telephone,
                'role': role,
            }

            if erreurs_ligne:
                membre['erreurs'] = erreurs_ligne
                membres_invalides.append(membre)
            else:
                membres_valides.append(membre)

        return membres_valides, membres_invalides


# ==========================================================
# AFFICHAGE / MODIFICATION D'UN MEMBRE
# ==========================================================

class MembreSerializer(serializers.ModelSerializer):

    # Date d'expiration du lien d'invitation (comptes en attente uniquement)
    invitation_expiration = serializers.DateTimeField(
        source='token_expiration', read_only=True
    )

    class Meta:

        model = Utilisateur

        fields = [
            'id',
            'prenom',
            'nom',
            'telephone',
            'email',
            'role',
            'fonction_bureau',
            'statut',
            'invitation_expiration',
        ]

        # Champs protégés :
        # telephone → identifiant de connexion, non modifiable depuis le profil
        # role   → modifiable uniquement via modifier_role (président)
        # statut → modifiable uniquement via suspendre/reactiver (président)
        # gie    → jamais modifiable par l'utilisateur (non exposé dans fields)
        read_only_fields = [
            'id',
            'telephone',
            'role',
            'fonction_bureau',
            'statut',
        ]

    def validate_prenom(self, valeur):
        return valider_nom_prenom(valeur, "prénom")

    def validate_nom(self, valeur):
        return valider_nom_prenom(valeur, "nom")

    def validate_telephone(self, valeur):
        utilisateur = self.instance
        telephone, erreurs = valider_telephone(
            valeur,
            exclure_id=utilisateur.id if utilisateur else None
        )
        if erreurs:
            raise serializers.ValidationError(erreurs)
        return telephone

    def validate_email(self, valeur):
        """
        Valide l'unicité de l'email en excluant l'utilisateur courant.
        Le champ est optionnel (null=True, blank=True dans le modèle).
        """
        if not valeur:
            return valeur

        valeur = str(valeur).strip().lower()
        utilisateur = self.instance

        if Utilisateur.objects.filter(email=valeur).exclude(id=utilisateur.id).exists():
            raise serializers.ValidationError(
                "Cette adresse email est déjà utilisée."
            )

        return valeur
